#!/usr/bin/env python3
"""Numera, valida, converte e inventaria laudos DOCX em lote."""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import platform
import re
import shutil
import subprocess
import sys
import unicodedata
import zipfile
from collections import Counter
from pathlib import Path
from typing import Iterable

from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}
DATE_RE = re.compile(
    r"Data\s+do\s+Recebimento\s*:\s*(\d{1,2}/\d{1,2}/\d{4})", re.IGNORECASE
)
TITLE_RE = re.compile(r"^Relat[oó]rio\s+de\s+Ensaio\s+N", re.IGNORECASE)
NUMBER_RE = re.compile(r"\d{5}\.\d{4}")
TEMP_PREFIX = "~$"


class AutomationError(RuntimeError):
    """Erro controlado, adequado para exibição no terminal."""


def natural_text(value: str) -> str:
    return unicodedata.normalize("NFKD", value).casefold()


def iter_docx(root: Path) -> tuple[list[Path], list[str]]:
    files: list[Path] = []
    excluded: list[str] = []
    for path in sorted(root.rglob("*.docx")):
        relative = path.relative_to(root).as_posix()
        if path.name.startswith(TEMP_PREFIX):
            excluded.append(relative)
        else:
            files.append(path)
    return files, excluded


def read_document_xml(path: Path) -> bytes:
    try:
        with zipfile.ZipFile(path) as archive:
            return archive.read("word/document.xml")
    except (zipfile.BadZipFile, KeyError) as exc:
        raise AutomationError(f"DOCX inválido: {path}") from exc


def paragraph_text(paragraph: etree._Element) -> str:
    return "".join(paragraph.xpath(".//w:t/text()", namespaces=NS))


def document_text(raw_xml: bytes) -> str:
    root = etree.fromstring(raw_xml)
    return "\n".join(
        paragraph_text(p) for p in root.xpath(".//w:p", namespaces=NS)
    )


def receipt_date(text: str, path: Path) -> dt.date:
    values = list(dict.fromkeys(DATE_RE.findall(text)))
    if len(values) != 1:
        raise AutomationError(
            f"Data do Recebimento ausente ou ambígua em '{path}': {values}"
        )
    try:
        return dt.datetime.strptime(values[0], "%d/%m/%Y").date()
    except ValueError as exc:
        raise AutomationError(f"Data inválida em '{path}': {values[0]}") from exc


def locate_title(root: etree._Element, path: Path) -> etree._Element:
    matches = [
        p
        for p in root.xpath(".//w:p", namespaces=NS)
        if TITLE_RE.match(paragraph_text(p).strip())
    ]
    if len(matches) != 1:
        raise AutomationError(
            f"Era esperado um único título 'Relatório de Ensaio' em '{path}'; "
            f"encontrados: {len(matches)}"
        )
    return matches[0]


def replace_visible_text(
    paragraph: etree._Element, start: int, end: int, replacement: str
) -> None:
    nodes = paragraph.xpath(".//w:t", namespaces=NS)
    cursor = 0
    inserted = False
    for node in nodes:
        original = node.text or ""
        node_start, node_end = cursor, cursor + len(original)
        cursor = node_end
        if node_end <= start or node_start >= end:
            continue
        local_start = max(0, start - node_start)
        local_end = min(len(original), end - node_start)
        prefix = original[:local_start]
        suffix = original[local_end:]
        if not inserted:
            node.text = prefix + replacement + suffix
            inserted = True
        else:
            node.text = prefix + suffix
    if not inserted:
        raise AutomationError("Não foi possível localizar o texto visível para substituição")


def update_title_number(raw_xml: bytes, report_number: str, path: Path) -> tuple[bytes, str, str]:
    parser = etree.XMLParser(remove_blank_text=False)
    root = etree.fromstring(raw_xml, parser)
    title = locate_title(root, path)
    before = paragraph_text(title)
    match = NUMBER_RE.search(before)
    if match:
        replace_visible_text(title, match.start(), match.end(), report_number)
    else:
        nodes = title.xpath(".//w:t", namespaces=NS)
        if not nodes:
            raise AutomationError(f"Título sem texto em '{path}'")
        separator = "" if before.endswith((" ", "\t", "\n")) else " "
        nodes[-1].text = (nodes[-1].text or "") + separator + report_number
    after = paragraph_text(title)
    if after.count(report_number) != 1:
        raise AutomationError(f"Falha ao inserir o número em '{path}'")
    updated = etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)
    return updated, before, after


def write_patched_docx(source: Path, target: Path, report_number: str) -> tuple[str, str]:
    target.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(source, "r") as src, zipfile.ZipFile(target, "w") as dst:
        before = after = ""
        for info in src.infolist():
            payload = src.read(info.filename)
            if info.filename == "word/document.xml":
                payload, before, after = update_title_number(
                    payload, report_number, source
                )
            dst.writestr(info, payload)
    if not after:
        raise AutomationError(f"document.xml não processado em '{source}'")
    return before, after


def inventory(input_root: Path, start: int, report_year: int) -> tuple[list[dict], list[str]]:
    files, excluded = iter_docx(input_root)
    if not files:
        raise AutomationError(f"Nenhum DOCX encontrado em '{input_root}'")
    rows: list[dict] = []
    for source in files:
        raw = read_document_xml(source)
        text = document_text(raw)
        date = receipt_date(text, source)
        rows.append(
            {
                "source": source,
                "relative_path": source.relative_to(input_root).as_posix(),
                "receipt_date": date,
            }
        )
    rows.sort(
        key=lambda row: (
            row["receipt_date"].month,
            row["receipt_date"].day,
            natural_text(row["relative_path"]),
            row["receipt_date"].year,
        )
    )
    for index, row in enumerate(rows):
        row["order"] = index + 1
        row["report_number"] = f"{start + index:05d}.{report_year}"
        row["outside_report_year"] = row["receipt_date"].year != report_year
    return rows, excluded


def clean_output(path: Path, force: bool) -> None:
    if not path.exists():
        return
    if any(path.iterdir()) and not force:
        raise AutomationError(
            f"A saída '{path}' não está vazia. Use --force para recriá-la."
        )
    if force:
        shutil.rmtree(path)


def number_documents(
    input_root: Path,
    output_root: Path,
    start: int,
    report_year: int,
    force: bool,
) -> tuple[list[dict], list[str]]:
    if input_root.resolve() == output_root.resolve():
        raise AutomationError("Entrada e saída devem ser pastas diferentes")
    clean_output(output_root, force)
    output_root.mkdir(parents=True, exist_ok=True)
    rows, excluded = inventory(input_root, start, report_year)
    for row in rows:
        source = row["source"]
        target = output_root / row["relative_path"]
        before, after = write_patched_docx(source, target, row["report_number"])
        row["output"] = target
        row["title_before"] = before
        row["title_after"] = after
    validate_numbered_set(input_root, output_root, rows)
    return rows, excluded


def normalized_without_title(raw_xml: bytes, path: Path) -> bytes:
    root = etree.fromstring(raw_xml)
    title = locate_title(root, path)
    parent = title.getparent()
    if parent is None:
        raise AutomationError(f"Título sem elemento pai em '{path}'")
    index = parent.index(title)
    placeholder = etree.Element("TITLE_PLACEHOLDER")
    parent.remove(title)
    parent.insert(index, placeholder)
    return etree.tostring(root, method="c14n")


def validate_numbered_set(input_root: Path, output_root: Path, rows: list[dict]) -> None:
    expected = {row["relative_path"] for row in rows}
    actual = {
        path.relative_to(output_root).as_posix()
        for path in output_root.rglob("*.docx")
        if not path.name.startswith(TEMP_PREFIX)
    }
    if actual != expected:
        raise AutomationError("A lista de DOCX da saída não corresponde à entrada")
    numbers = [row["report_number"] for row in rows]
    if len(numbers) != len(set(numbers)):
        raise AutomationError("A validação encontrou números repetidos")
    for row in rows:
        source = input_root / row["relative_path"]
        target = output_root / row["relative_path"]
        with zipfile.ZipFile(source) as src, zipfile.ZipFile(target) as dst:
            if src.namelist() != dst.namelist():
                raise AutomationError(f"Estrutura OOXML alterada em '{row['relative_path']}'")
            for member in src.namelist():
                if member != "word/document.xml" and src.read(member) != dst.read(member):
                    raise AutomationError(
                        f"Conteúdo fora do texto principal alterado: {row['relative_path']}"
                    )
            src_xml = src.read("word/document.xml")
            dst_xml = dst.read("word/document.xml")
        if normalized_without_title(src_xml, source) != normalized_without_title(dst_xml, target):
            raise AutomationError(f"Conteúdo fora do título alterado: {row['relative_path']}")
        title = paragraph_text(locate_title(etree.fromstring(dst_xml), target))
        if title.count(row["report_number"]) != 1:
            raise AutomationError(f"Número ausente ou duplicado em '{row['relative_path']}'")


def autosize_sheet(sheet) -> None:
    for column in range(1, sheet.max_column + 1):
        width = max(
            len(str(sheet.cell(row=row, column=column).value or ""))
            for row in range(1, sheet.max_row + 1)
        )
        sheet.column_dimensions[get_column_letter(column)].width = min(width + 2, 70)


def add_rows_sheet(workbook: Workbook, name: str, rows: list[dict]) -> None:
    sheet = workbook.create_sheet(name)
    headers = ["Ordem", "Número", "Data do recebimento", "Ano divergente", "Pasta/arquivo"]
    sheet.append(headers)
    for row in rows:
        sheet.append(
            [
                row["order"],
                row["report_number"],
                row["receipt_date"].strftime("%d/%m/%Y"),
                "SIM" if row["outside_report_year"] else "NÃO",
                row["relative_path"],
            ]
        )
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    autosize_sheet(sheet)


def write_control_files(rows: list[dict], excluded: list[str], destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_rows_sheet(workbook, "Controle", rows)
    divergent = [row for row in rows if row["outside_report_year"]]
    add_rows_sheet(workbook, "Datas fora do ano", divergent)
    summary = workbook.create_sheet("Resumo", 0)
    summary.append(["Indicador", "Valor"])
    summary.append(["Total de documentos", len(rows)])
    summary.append(["Primeiro número", rows[0]["report_number"]])
    summary.append(["Último número", rows[-1]["report_number"]])
    summary.append(["Datas fora do ano do relatório", len(divergent)])
    summary.append(["Temporários ignorados", len(excluded)])
    summary.append(["Documentos por pasta", json.dumps(dict(Counter(Path(r["relative_path"]).parts[0] for r in rows)), ensure_ascii=False)])
    for cell in summary[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
    autosize_sheet(summary)
    workbook.save(destination / "controle_de_laudos.xlsx")

    with (destination / "controle_de_laudos.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(["ordem", "numero", "data_recebimento", "ano_divergente", "arquivo"])
        for row in rows:
            writer.writerow(
                [
                    row["order"],
                    row["report_number"],
                    row["receipt_date"].strftime("%d/%m/%Y"),
                    "SIM" if row["outside_report_year"] else "NÃO",
                    row["relative_path"],
                ]
            )


def find_soffice() -> str | None:
    candidates = [
        shutil.which("soffice"),
        shutil.which("libreoffice"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    return next((str(item) for item in candidates if item and Path(item).exists()), None)


def convert_with_libreoffice(docx_files: Iterable[Path], input_root: Path, pdf_root: Path) -> int:
    executable = find_soffice()
    if not executable:
        raise AutomationError("LibreOffice não encontrado")
    count = 0
    for source in docx_files:
        target_dir = pdf_root / source.relative_to(input_root).parent
        target_dir.mkdir(parents=True, exist_ok=True)
        result = subprocess.run(
            [executable, "--headless", "--convert-to", "pdf", "--outdir", str(target_dir), str(source)],
            capture_output=True,
            text=True,
        )
        target = target_dir / f"{source.stem}.pdf"
        if result.returncode != 0 or not target.exists() or target.stat().st_size == 0:
            raise AutomationError(
                f"Falha ao converter '{source}'. Saída: {result.stdout} {result.stderr}"
            )
        count += 1
    return count


def convert_with_word(docx_files: Iterable[Path], input_root: Path, pdf_root: Path) -> int:
    try:
        import win32com.client  # type: ignore
    except ImportError as exc:
        raise AutomationError("pywin32 não está instalado") from exc
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    count = 0
    try:
        for source in docx_files:
            target = pdf_root / source.relative_to(input_root).with_suffix(".pdf")
            target.parent.mkdir(parents=True, exist_ok=True)
            document = word.Documents.Open(str(source.resolve()), ReadOnly=True)
            try:
                document.ExportAsFixedFormat(str(target.resolve()), 17)
            finally:
                document.Close(False)
            if not target.exists() or target.stat().st_size == 0:
                raise AutomationError(f"PDF não foi criado para '{source}'")
            count += 1
    finally:
        word.Quit()
    return count


def convert_documents(input_root: Path, pdf_root: Path, engine: str, force: bool) -> int:
    files, _ = iter_docx(input_root)
    if not files:
        raise AutomationError(f"Nenhum DOCX encontrado em '{input_root}'")
    clean_output(pdf_root, force)
    pdf_root.mkdir(parents=True, exist_ok=True)
    chosen = engine
    if chosen == "auto":
        chosen = "word" if platform.system() == "Windows" else "libreoffice"
    if chosen == "word":
        return convert_with_word(files, input_root, pdf_root)
    return convert_with_libreoffice(files, input_root, pdf_root)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Numera, valida, converte e inventaria laudos DOCX em lote."
    )
    sub = parser.add_subparsers(dest="command", required=True)

    number = sub.add_parser("numerar", help="Numera os DOCX e gera a planilha")
    number.add_argument("entrada", type=Path)
    number.add_argument("saida", type=Path)
    number.add_argument("--inicio", type=int, required=True, help="Primeiro número, sem zeros")
    number.add_argument("--ano", type=int, required=True, help="Ano do número do relatório")
    number.add_argument("--force", action="store_true", help="Recria a pasta de saída")

    convert = sub.add_parser("converter", help="Converte DOCX em PDF")
    convert.add_argument("entrada", type=Path)
    convert.add_argument("saida", type=Path)
    convert.add_argument("--motor", choices=["auto", "word", "libreoffice"], default="auto")
    convert.add_argument("--force", action="store_true")

    run = sub.add_parser("executar", help="Numera, valida, gera planilha e converte")
    run.add_argument("entrada", type=Path)
    run.add_argument("saida", type=Path)
    run.add_argument("--inicio", type=int, required=True)
    run.add_argument("--ano", type=int, required=True)
    run.add_argument("--motor", choices=["auto", "word", "libreoffice"], default="auto")
    run.add_argument("--force", action="store_true")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    try:
        if args.command in {"numerar", "executar"}:
            numbered_root = args.saida / "docx_numerados" if args.command == "executar" else args.saida / "docx_numerados"
            rows, excluded = number_documents(
                args.entrada, numbered_root, args.inicio, args.ano, args.force
            )
            write_control_files(rows, excluded, args.saida)
            print(
                f"OK: {len(rows)} DOCX numerados de {rows[0]['report_number']} "
                f"a {rows[-1]['report_number']}."
            )
            if excluded:
                print(f"Temporários ignorados: {len(excluded)}")
            if args.command == "executar":
                count = convert_documents(
                    numbered_root, args.saida / "pdf", args.motor, args.force
                )
                print(f"OK: {count} PDFs gerados.")
        else:
            count = convert_documents(args.entrada, args.saida, args.motor, args.force)
            print(f"OK: {count} PDFs gerados.")
        return 0
    except AutomationError as exc:
        print(f"ERRO: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

